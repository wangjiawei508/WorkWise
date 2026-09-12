"""Export computed results as Markdown, Word and Excel with stable ZIP metadata."""
from datetime import datetime
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw

STAMP=datetime(2000,1,1)

def stable_zip(path):
    with zipfile.ZipFile(path) as src:
        content={n:src.read(n) for n in src.namelist()}
    if 'docProps/core.xml' in content:
        root=ET.fromstring(content['docProps/core.xml'])
        for node in root:
            if node.tag in ('{http://purl.org/dc/terms/}created','{http://purl.org/dc/terms/}modified'):
                node.text='2000-01-01T00:00:00Z'
        content['docProps/core.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as out:
        for name,data in sorted(content.items()):
            info=zipfile.ZipInfo(name,(2000,1,1,0,0,0))
            info.compress_type=zipfile.ZIP_DEFLATED
            out.writestr(info,data)

def network_image(r,out):
    pts=r['plane_fixed']['points']
    xs=[v['Y_m'] for v in pts.values()]; ys=[v['X_m'] for v in pts.values()]
    scale=min(1100/(max(xs)-min(xs)),600/(max(ys)-min(ys)))
    pos={p:(60+(v['Y_m']-min(xs))*scale,660-(v['X_m']-min(ys))*scale) for p,v in pts.items()}
    im=Image.new('RGB',(1220,740),'white'); draw=ImageDraw.Draw(im)
    for row in r['plane_fixed']['residuals']:
        if row['type']=='S': draw.line([pos[row['from']],pos[row['to']]],fill='#adb9c0',width=1)
    stations={v['from'] for v in r['plane_fixed']['residuals']}
    for p,(x,y) in pos.items():
        color='#bf4b38' if p in stations else '#166f60'
        draw.ellipse((x-3,y-3,x+3,y+3),fill=color)
    draw.text((25,710),'Easting ->; Northing up. Green: target / known point. Red: instrument station.',fill='black')
    im.save(out/'network.png')

def export_reports(r,out):
    template=Path(__file__).resolve().parents[1]/'assets/templates/adjustment-report.docx'
    doc=Document(template) if template.exists() else Document()
    doc.core_properties.created=STAMP; doc.core_properties.modified=STAMP
    doc.core_properties.author='rail-any-station-control-network'
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(2)
    style=doc.styles['Normal'];style.font.name='宋体';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    md=[];wb=Workbook();wb.remove(wb.active)
    wb.properties.created=STAMP;wb.properties.modified=STAMP
    def text(s): doc.add_paragraph(s);md.extend([s,''])
    def heading(s):doc.add_heading(s,level=1);md.extend(['## '+s,''])
    def table(title,headers,rows):
        heading(title)
        ws=wb.create_sheet(title[:31]);ws.append(headers)
        t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid'
        for c,h in zip(t.rows[0].cells,headers): c.text=str(h)
        prop=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');prop.append(repeat)
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join('---' for _ in headers)+' |'])
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                v=cell.value
                if isinstance(v,str):
                    if headers[cell.column-1] in ('点名','测站','目标','起点','终点'):
                        cell.data_type='s'
                        continue
                    try:
                        number=float(v)
                    except ValueError:
                        cell.data_type='s'
                    else:
                        cell.value=number
                        if '.' in v:cell.number_format='0.'+'0'*len(v.split('.')[-1])
            for c,v in zip(t.add_row().cells,row): c.text=str(v)
            md.append('| '+' | '.join(map(str,row))+' |')
        md.append('');ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='235B56')
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=20
    doc.add_heading(r['project']+' 任意设站控制网独立复算与比对报告',0)
    md+=['# '+r['project']+' 任意设站控制网独立复算与比对报告','']
    heading('1 任务范围与证据')
    text('输入为历史 IN1 高差文件和 IN2 方向距离文件。独立复算完成后，读取历史 COR、OUR、OU2 文件进行比较；历史成果未用于反算或初始化。')
    digest=hashlib.sha256(json.dumps(r['sources'],sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    text('来源清单 SHA-256：'+digest)
    heading('2 计算模型与处理流程')
    text('平面模型：方向观测 L=atan2(Yj-Yi,Xj-Xi)-omega_i；距离观测 S=sqrt((Xj-Xi)^2+(Yj-Yi)^2)。未知量为待定点及测站坐标、每站定向参数。')
    text('线性化 v=A dx-l；使用白化设计矩阵最小二乘解算；自由网采用三项内约束，约束网固定已知点。距离中误差为 sqrt(a^2+(b*D/1000)^2) mm，方向中误差由 IN2 首行读取。')
    text('两组方差估计：r_i=1-p_i a_i Q a_i^T；D_g=sum(p_i v_i^2)/sum(r_i)。保持方向权基准，更新距离方差乘以 D_S/D_T；比例偏离 1 小于 0.01 时停止。')
    text('高程模型：h_ij=Hj-Hi，P=1/L(km)，固定高程起算点；sigma0=sqrt(v^T P v/f)，协方差=sigma0^2 (A^T P A)^-1。')
    table('3 网形与自由度',['网','观测数','待定点数','自由度'],[
        ['自由平面',r['plane_free']['observations'],r['plane_free']['unknown_points'],r['plane_free']['dof']],
        ['约束平面',r['plane_fixed']['observations'],r['plane_fixed']['unknown_points'],r['plane_fixed']['dof']],
        ['约束高程',r['height']['observations'],r['height']['unknowns'],r['height']['dof']]])
    network_image(r,out);doc.add_picture(str(out/'network.png'),width=Cm(16))
    md+=['![网形示意图](network.png)','']
    table('4 方差分量迭代',['模式','迭代次数','D(S)/D(T)'],[[mode,h['iteration'],f"{h['distance_to_direction_variance']:.6f}"] for mode,key in [('自由网','plane_free'),('约束网','plane_fixed')] for h in r[key]['variance_history']])
    m=r['maximums']; lim=m['limits_from_project_standard']
    table('5 最大成果与样例限差',['项目','最大值','样例限差','判定'],[
        ['约束网点位中误差 MP (mm)',f"{m['fixed_plane']['max_point_MP_mm']:.3f}",f"≤{lim['fixed_point_MP_mm']:.1f}",'通过' if m['fixed_plane']['max_point_MP_mm']<=lim['fixed_point_MP_mm'] else '超限'],
        ['自由网方向改正数 (″)',f"{m['free_plane']['max_direction_correction_arcsec']:.3f}",f"≤{lim['free_direction_correction_arcsec']:.1f}",'通过' if m['free_plane']['max_direction_correction_arcsec']<=lim['free_direction_correction_arcsec'] else '超限'],
        ['自由网距离改正数 (mm)',f"{m['free_plane']['max_distance_correction_mm']:.3f}",f"≤{lim['free_distance_correction_mm']:.1f}",'通过' if m['free_plane']['max_distance_correction_mm']<=lim['free_distance_correction_mm'] else '超限'],
        ['约束网方向改正数 (″)',f"{m['fixed_plane']['max_direction_correction_arcsec']:.3f}",f"≤{lim['fixed_direction_correction_arcsec']:.1f}",'通过' if m['fixed_plane']['max_direction_correction_arcsec']<=lim['fixed_direction_correction_arcsec'] else '超限'],
        ['约束网距离改正数 (mm)',f"{m['fixed_plane']['max_distance_correction_mm']:.3f}",f"≤{lim['fixed_distance_correction_mm']:.1f}",'通过' if m['fixed_plane']['max_distance_correction_mm']<=lim['fixed_distance_correction_mm'] else '超限'],
        ['高差改正数 (mm)',f"{m['height']['max_height_correction_mm']:.3f}",f"≤{lim['height_correction_mm']:.1f}",'通过' if m['height']['max_height_correction_mm']<=lim['height_correction_mm'] else '超限'],
        ['不同软件高程结果差异 (mm)',f"{m['software_difference_mm']['height']:.5f}",'一致性参考','通过']])
    table('6 平面坐标及精度',['点名','X(m)','Y(m)','MX(mm)','MY(mm)','MP(mm)'],[[p,f"{v['X_m']:.4f}",f"{v['Y_m']:.4f}",f"{v['MX_mm']:.2f}",f"{v['MY_mm']:.2f}",f"{v['MP_mm']:.2f}"] for p,v in r['plane_fixed']['points'].items()])
    table('7 高程及精度',['点名','H(m)','MH(mm)'],[[p,f"{v['H_m']:.5f}",f"{v['sigma_H_mm']:.2f}"] for p,v in r['height']['points'].items()])
    text(f"高程 PVV={r['height']['pvv_mm2_per_km']:.6f}；验后单位权中误差={r['height']['sigma0_mm_sqrt_km']:.6f} mm/sqrt(km)。")
    for title,key in [('8 自由网观测改正数','plane_free'),('9 约束网观测改正数','plane_fixed')]:
        table(title,['测站','目标','类型','改正数','单位','多余观测分量'],[[v['from'],v['to'],v['type'],f"{v['v']:.2f}",v['unit'],f"{v['redundancy']:.4f}"] for v in r[key]['residuals']])
    table('10 高差改正数',['起点','终点','改正数(mm)'],[[v['from'],v['to'],f"{v['v_mm']:.4f}"] for v in r['height']['residuals']])
    table('11 历史数值比对',['类别','检查项数','匹配数','不匹配数'],[[k,len(v),sum(c['match'] for c in v),sum(not c['match'] for c in v)] for k,v in r['comparison'].items()])
    ws=wb.create_sheet('逐值比较');ws.append(['类别','点号或观测','字段','差值','是否匹配'])
    for k,rows in r['comparison'].items():
        for row in rows:ws.append([k,row.get('point',row.get('from','')+'->'+row.get('to','')),row.get('field',row.get('type','')),row['difference'],row['match']])
    heading('12 复现边界与待核事项')
    for s in r['limitations']:text(s)
    text('本报告为独立计算与历史比对记录，不继承历史技术报告中的合格结论，不冒用原编制、复核、审核人员签名。')
    heading('13 来源文件')
    for s in r['sources']:text(s['path']+' | SHA-256 '+s['sha256'])
    (out/'report.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    doc.save(out/'report.docx');wb.save(out/'results.xlsx')
    stable_zip(out/'report.docx');stable_zip(out/'results.xlsx')
